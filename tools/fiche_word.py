#!/usr/bin/env python3
"""Fiche de validation au format Word, pour l'auteur du cahier (qui n'utilise ni Claude ni Copilot).

  export [projet] [cahier.docx]  → projets/<p>/reports/validation_<cahier>.docx : partie 1 = questions
                                   ouvertes (V1), partie 2 = scénarios reformulés à approuver (V2).
                                   Sans cahier : un fichier par cahier cité dans les .feature du projet.
  import <fiche.docx> [projet]   relit la fiche remplie : chaque réponse non vide est recopiée sur la ligne
                                   « réponse: » de la question dans le .feature ; les avis « Conforme » (V2)
                                   sont affichés. Les statuts ne changent pas ici : c'est le rôle de /valider.
[projet] : facultatif s'il n'y a qu'un projet, ou si la fiche est dans projets/<p>/.
"""
from __future__ import annotations

import re
import sys
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import projets as P  # noqa: E402
from valider_features import Scenario, charger_tous  # noqa: E402

ENTETE_QUESTIONS = ["Réf.", "Étape", "Question", "Proposition par défaut", "Votre réponse"]
CONFORME_RE = re.compile(r"^([A-Z][A-Z0-9]*-\d+) — conforme \?", re.IGNORECASE)


def cahier(sc: Scenario) -> str:
    return sc.meta.get("source", "?").split("§")[0].strip()


def _tableau(doc, lignes: list[list[str]], largeurs: list[float] | None = None):
    table = doc.add_table(rows=len(lignes), cols=len(lignes[0]))
    table.style = "Table Grid"
    for i, ligne in enumerate(lignes):
        for j, valeur in enumerate(ligne):
            cellule = table.cell(i, j)
            cellule.text = valeur
            if i == 0:
                for run in cellule.paragraphs[0].runs:
                    run.bold = True
    return table


def exporter(scenarios: list[Scenario], source: str, projet: Path) -> Path:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(f"Validation du cahier « {source} »", 0)
    doc.add_paragraph(
        "Partie 1 : répondez dans la dernière colonne. Une case laissée vide signifie que la proposition "
        "par défaut vous convient. Les questions marquées BLOQUANT doivent recevoir une réponse.\n"
        "Partie 2 : relisez chaque scénario reformulé et indiquez Oui ou Non dans « Conforme ? ».\n"
        "Ne modifiez pas la colonne « Réf. » : elle sert à relire vos réponses automatiquement.")

    doc.add_heading("Partie 1 — Questions", 1)
    for sc in scenarios:
        ouvertes = sorted((q for q in sc.questions if not q["reponse"]), key=lambda q: not q["bloquant"])
        doc.add_heading(f"{sc.id} — {sc.titre}", 2)
        if not ouvertes:
            doc.add_paragraph("Aucune question ouverte.")
            continue
        lignes = [ENTETE_QUESTIONS]
        for q in ouvertes:
            etape = "tout le scénario" if q["etape"] == "-" else next(
                (e.texte for e in sc.etapes if str(e.n) == q["etape"]), q["etape"])
            question = ("BLOQUANT — " if q["bloquant"] else "") + q["question"]
            lignes.append([f"{sc.id} {q['id']}", etape, question, q["proposition"], ""])
        _tableau(doc, lignes)

    doc.add_heading("Partie 2 — Est-ce bien ce que vous vouliez tester ?", 1)
    for sc in scenarios:
        doc.add_heading(f"{sc.id} — {sc.titre}", 2)
        if sc.meta.get("objectif"):
            p = doc.add_paragraph(sc.meta["objectif"])
            p.runs[0].italic = True
        for e in sc.etapes:
            p = doc.add_paragraph(e.texte, style="List Number")
            if e.confiance == "basse":
                r = p.add_run("  (interprétation à vérifier)")
                r.font.color.rgb = RGBColor(0xC0, 0x50, 0x00)
        _tableau(doc, [[f"{sc.id} — Conforme ? (Oui / Non)", "Remarques"], ["", ""]])

    sortie = projet / "reports" / f"validation_{Path(source).stem}.docx"
    sortie.parent.mkdir(exist_ok=True)
    doc.save(sortie)
    return sortie


def importer(fiche: Path, projet: Path) -> int:
    doc = Document(fiche)
    reponses: dict[str, dict[str, str]] = defaultdict(dict)
    avis = []
    for table in doc.tables:
        entete = [c.text.strip() for c in table.rows[0].cells]
        if entete[:1] == ENTETE_QUESTIONS[:1] and len(entete) >= 5:
            for row in table.rows[1:]:
                cellules = [c.text.strip() for c in row.cells]
                m = re.fullmatch(r"([A-Z][A-Z0-9]*-\d+)\s+(Q\d+)", cellules[0])
                if m and cellules[4]:
                    reponses[m[1]][m[2]] = " ".join(cellules[4].split())
        elif CONFORME_RE.match(entete[0]) and len(table.rows) > 1:
            valeur, remarque = (c.text.strip() for c in table.rows[1].cells[:2])
            avis.append((CONFORME_RE.match(entete[0])[1], valeur or "(vide)", remarque))

    total = 0
    for sc_id, par_question in reponses.items():
        fichiers = sorted((projet / "features").glob(f"{sc_id}_*.feature"))
        if not fichiers:
            print(f"✗ {sc_id} : aucun .feature dans {P.relatif(projet / 'features')}")
            continue
        texte = fichiers[0].read_text(encoding="utf-8")
        for q_id, reponse in par_question.items():
            motif = re.compile(rf"(^# {q_id} \|.*\n(?:#\s+(?!r[ée]ponse).*\n)*#\s+r[ée]ponse\s*:)[^\n]*", re.MULTILINE)
            texte, nb = motif.subn(lambda m: f"{m[1]} {reponse}", texte, count=1)
            print(f"{'✓' if nb else '✗'} {sc_id} {q_id} : {reponse if nb else 'question introuvable'}")
            total += nb
        fichiers[0].write_text(texte, encoding="utf-8")
    print(f"\n{total} réponse(s) reportée(s) dans les .feature.")
    if avis:
        print("\nAvis de l'auteur (V2) :")
        for sc_id, valeur, remarque in avis:
            print(f"  {sc_id} : conforme = {valeur}" + (f" — {remarque}" if remarque else ""))
    print("\nÉtape suivante : /valider (statuts, questions sans réponse = proposition par défaut si accord).")
    return 0


def main(argv: list[str]) -> int:
    for _flux in (sys.stdout, sys.stderr):
        _flux.reconfigure(encoding="utf-8")
    if not argv or argv[0] not in {"export", "import"}:
        print(__doc__)
        return 2
    if argv[0] == "import":
        if len(argv) < 2:
            print("Indiquer la fiche remplie : import projets/<projet>/reports/validation_<cahier>.docx")
            return 2
        projet = P.depuis_chemin(argv[1]) or P.separer(argv[2:])[0]
        return importer(Path(argv[1]), projet)
    projet, reste = P.separer(argv[1:])
    par_cahier: dict[str, list[Scenario]] = defaultdict(list)
    for sc, _ in charger_tous(projet):
        if sc.etapes:
            par_cahier[cahier(sc)].append(sc)
    cibles = [Path(reste[0]).name] if reste else sorted(par_cahier)
    for source in cibles:
        if source not in par_cahier:
            print(f"✗ aucun scénario ne cite « {source} » (sources connues : {', '.join(par_cahier) or 'aucune'})")
            return 1
        print(f"→ {P.relatif(exporter(par_cahier[source], source, projet))}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
